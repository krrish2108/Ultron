import copy
import re
import shutil
import time
import uuid
import asyncio
from typing import cast
from pathlib import Path

import pypandoc
from docx import Document as DocxDocument
from langchain_core.messages import HumanMessage, SystemMessage
from langchain_ollama import ChatOllama

from services.utils.config import DOC_GEN_MODEL
from services.agents.models import DocGenState, PlaceholderValues, BlockContent


PLACEHOLDER_PATTERN = re.compile(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}")
BLOCK_PATTERN = re.compile(r"\{\{\s*block:\s*([a-zA-Z0-9_]+)\s*\}\}")


PLACEHOLDER_FILL_SYSTEM_PROMPT = """\
You are a document-filling assistant. You are given a list of PLACEHOLDER NAMES \
that appear in a Word template, and RAW SOURCE CONTENT to pull facts from.

Your job is to return a value for every placeholder name, using information found \
in the raw source content.

## Rules

1. Return a value for every placeholder name provided — do not skip any.
2. Use ONLY information present in the raw source content. Do not invent facts.
3. If nothing relevant is found in the source content for a given placeholder, \
return an empty string "" for that placeholder — do NOT guess or fabricate.
4. Keep each value concise and appropriate for inserting directly into a sentence, \
table cell, or field (e.g. a name, a date, a number, a short phrase). Do not include \
the placeholder's own name or braces in the value.
5. Do not add commentary, explanations, or extra placeholders not in the list.
"""

BLOCK_FILL_SYSTEM_PROMPT = """\
You are a document-drafting assistant. You are given a list of SECTION NAMES that \
correspond to content blocks inside a Word template, and RAW SOURCE CONTENT to draft from.

Each section name is marked with a PLACEMENT TYPE that tells you the shape the \
content must take:

- **standalone**: this marker sits alone on its own line in the template. You may \
use one or more paragraphs, a bullet list, or a markdown table — whatever best fits \
the source content for that section.
- **inline**: this marker sits in the middle of a sentence alongside other text \
(e.g. "Status: {{block:status_note}}."). You MUST write a single short plain-text \
phrase or sentence for this one — no headings, no bullet lists, no markdown tables, \
no line breaks. It has to read naturally as part of the surrounding sentence.

## Rules

1. Return content for every section name given — do not skip any.
2. Base the content only on the RAW SOURCE CONTENT provided. Do not invent facts.
3. If nothing relevant exists in the source content for a section, return a single \
short sentence (or, for inline sections, a short phrase) noting that no information \
was available, rather than inventing content.
4. Do NOT include the section name or a heading in your output — the template already \
has the heading in place (for standalone sections); provide only the body content \
that follows it.
5. For standalone sections, use proper markdown table syntax (pipes, header separator \
row) if a table fits best.
6. For inline sections, never use markdown syntax (no #, *, -, |, or line breaks) — \
plain text only.
7. Do not add commentary outside of the requested section content.
"""


def _iter_paragraphs(doc):  # type: ignore[type-arg]
    """Yield all paragraphs in the document body AND inside table cells."""
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def read_template_node(state: DocGenState) -> dict:
    """Node: Inspect the uploaded template (if any)."""
    template_path = state.get("template_path")
    empty = {
        "template_outline": None,
        "template_placeholders": None,
        "template_blocks": None,
        "template_block_placements": None,
    }
    if not template_path:
        print("[DOC_GEN] read_template: no template supplied, skipping")
        return empty

    template_file = Path(template_path)
    if not template_file.exists() or template_file.suffix.lower() != ".docx":
        print(f"[DOC_GEN] read_template: invalid template_path ({template_path}), skipping")
        return empty

    try:
        doc = DocxDocument(str(template_file))
        simple_found: set[str] = set()
        block_found: set[str] = set()
        block_placements: dict[str, str] = {}
        for p in _iter_paragraphs(doc):
            full_text = "".join(run.text for run in p.runs)
            for match in BLOCK_PATTERN.finditer(full_text):
                name = match.group(1)
                block_found.add(name)
                is_standalone = bool(BLOCK_PATTERN.fullmatch(full_text.strip()))
                placement = "standalone" if is_standalone else "inline"
                if block_placements.get(name) != "inline":
                    block_placements[name] = placement
            for match in PLACEHOLDER_PATTERN.finditer(full_text):
                simple_found.add(match.group(1))

        if simple_found or block_found:
            print(f"[DOC_GEN] read_template: found {len(simple_found)} placeholder(s) "
                  f"{sorted(simple_found)} and {len(block_found)} block(s) {block_placements}")
            return {
                "template_outline": None,
                "template_placeholders": sorted(simple_found) if simple_found else None,
                "template_blocks": sorted(block_found) if block_found else None,
                "template_block_placements": block_placements if block_placements else None,
            }
    except Exception as e:
        print(f"[DOC_GEN] read_template: python-docx FAILED to scan template — {e}")

    try:
        template_md = pypandoc.convert_file(str(template_file), "md")
    except Exception as e:
        print(f"[DOC_GEN] read_template: pypandoc FAILED to read template — {e}")
        return empty

    heading_lines = [
        line.strip() for line in template_md.splitlines()  # type: ignore
        if re.match(r"^#{1,6}\s+\S", line.strip())  # type: ignore
    ]

    if not heading_lines:
        print("[DOC_GEN] read_template: no placeholders, blocks, or headings found, skipping")
        return empty

    outline = "\n".join(heading_lines)
    print(f"[DOC_GEN] read_template: extracted {len(heading_lines)} heading(s) from template")
    return {
        "template_outline": outline,
        "template_placeholders": None,
        "template_blocks": None,
        "template_block_placements": None,
    }


async def fill_placeholders_node(state: DocGenState) -> dict:
    placeholders = state.get("template_placeholders") or []
    if not placeholders:
        print("[DOC_GEN] fill_placeholders: no simple placeholders to fill, skipping")
        return {"placeholder_values": {}}

    context_md = state.get("context_md", "")
    message = state.get("message", "")

    llm = ChatOllama(model=DOC_GEN_MODEL, num_ctx=8192, temperature=0.2)
    structured_llm = llm.with_structured_output(PlaceholderValues)

    print(f"[DOC_GEN] fill_placeholders: filling {len(placeholders)} placeholder(s)")
    t0 = time.monotonic()
    try:
        result = cast(PlaceholderValues, await asyncio.wait_for(
            structured_llm.ainvoke([
                SystemMessage(content=PLACEHOLDER_FILL_SYSTEM_PROMPT),
                HumanMessage(content=(
                    f"User request: {message}\n\n"
                    f"PLACEHOLDER NAMES to fill:\n{', '.join(placeholders)}\n\n"
                    f"RAW SOURCE CONTENT:\n{context_md}"
                )),
            ]),
            timeout=180,
        ))
        values = result.values
    except asyncio.TimeoutError:
        print("[DOC_GEN] fill_placeholders: TIMED OUT")
        values = {}
    except Exception as e:
        print(f"[DOC_GEN] fill_placeholders: FAILED — {type(e).__name__}: {e}")
        values = {}

    for name in placeholders:
        values.setdefault(name, "")

    elapsed = time.monotonic() - t0
    print(f"[DOC_GEN] fill_placeholders: got {len(values)} value(s) in {elapsed:.1f}s")
    return {"placeholder_values": values}


async def fill_blocks_node(state: DocGenState) -> dict:
    block_names = state.get("template_blocks") or []
    if not block_names:
        print("[DOC_GEN] fill_blocks: no block sections to fill, skipping")
        return {"block_content": {}}

    context_md = state.get("context_md", "")
    message = state.get("message", "")
    placements = state.get("template_block_placements") or {}

    section_lines = []
    for name in block_names:
        placement = placements.get(name, "standalone")
        if placement == "inline":
            hint = "inline — MUST be a single short plain-text phrase/sentence, no markdown"
        else:
            hint = "standalone — may use paragraphs, a bullet list, or a table"
        section_lines.append(f"- {name} ({hint})")
    section_list_text = "\n".join(section_lines)

    llm = ChatOllama(model=DOC_GEN_MODEL, num_ctx=8192, temperature=0.2)
    structured_llm = llm.with_structured_output(BlockContent)

    print(f"[DOC_GEN] fill_blocks: drafting {len(block_names)} block section(s) "
          f"({sum(1 for v in placements.values() if v == 'inline')} inline)")
    t0 = time.monotonic()
    try:
        result = cast(BlockContent, await asyncio.wait_for(
            structured_llm.ainvoke([
                SystemMessage(content=BLOCK_FILL_SYSTEM_PROMPT),
                HumanMessage(content=(
                    f"User request: {message}\n\n"
                    f"SECTION NAMES to draft (with placement type):\n{section_list_text}\n\n"
                    f"RAW SOURCE CONTENT:\n{context_md}"
                )),
            ]),
            timeout=240,
        ))
        blocks = result.blocks
    except asyncio.TimeoutError:
        print("[DOC_GEN] fill_blocks: TIMED OUT")
        blocks = {}
    except Exception as e:
        print(f"[DOC_GEN] fill_blocks: FAILED — {type(e).__name__}: {e}")
        blocks = {}

    for name in block_names:
        blocks.setdefault(name, "")

    elapsed = time.monotonic() - t0
    print(f"[DOC_GEN] fill_blocks: got {len(blocks)} block(s) in {elapsed:.1f}s")
    return {"block_content": blocks}


def _replace_placeholders_in_paragraph(paragraph, values: dict[str, str]) -> None:
    full_text = "".join(run.text for run in paragraph.runs)
    if not PLACEHOLDER_PATTERN.search(full_text):
        return

    def _sub(match: re.Match) -> str:
        name = match.group(1)
        return str(values.get(name, match.group(0)))

    new_text = PLACEHOLDER_PATTERN.sub(_sub, full_text)

    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)


def _insert_block_content_after(paragraph, markdown_text: str, tmp_dir: Path) -> None:
    tmp_docx = tmp_dir / f"block_{uuid.uuid4().hex}.docx"
    try:
        pypandoc.convert_text(markdown_text, "docx", format="md", outputfile=str(tmp_docx))
        sub_doc = DocxDocument(str(tmp_docx))

        anchor = paragraph._p
        for child in sub_doc.element.body:
            if child.tag.endswith('}sectPr'):
                continue
            new_child = copy.deepcopy(child)
            anchor.addnext(new_child)
            anchor = new_child
    finally:
        tmp_docx.unlink(missing_ok=True)


def _flatten_markdown_to_inline_text(markdown_text: str) -> str:
    if not markdown_text.strip():
        return ""
    try:
        plain = pypandoc.convert_text(markdown_text, "plain", format="md") # type: ignore
    except Exception as e:
        print(f"[DOC_GEN] _flatten_markdown_to_inline_text: pypandoc FAILED, using fallback strip — {e}")
        plain = re.sub(r'[#*_`>-]', '', markdown_text)
    return " ".join(str(plain).split())


def _replace_inline_blocks_in_paragraph(paragraph, block_content: dict[str, str]) -> None:
    full_text = "".join(run.text for run in paragraph.runs)
    if not BLOCK_PATTERN.search(full_text):
        return

    def _sub(match: re.Match) -> str:
        name = match.group(1)
        return _flatten_markdown_to_inline_text(block_content.get(name, ""))

    new_text = BLOCK_PATTERN.sub(_sub, full_text)

    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)


def render_placeholder_docx_node(state: DocGenState) -> dict:
    template_path: str = state["template_path"]  # type: ignore[assignment]
    values = state.get("placeholder_values") or {}
    block_content = state.get("block_content") or {}

    output_path = str(Path(__file__).resolve().parent.parent.parent / f"output_{uuid.uuid1()}.docx")
    shutil.copy(template_path, output_path)
    output_dir = Path(output_path).parent

    try:
        doc = DocxDocument(output_path)

        for p in _iter_paragraphs(doc):
            _replace_placeholders_in_paragraph(p, values)

        standalone_blocks: list[tuple] = []
        inline_paragraphs: list = []
        for p in _iter_paragraphs(doc):
            full_text = "".join(run.text for run in p.runs)
            match = BLOCK_PATTERN.search(full_text)
            if not match:
                continue
            if BLOCK_PATTERN.fullmatch(full_text.strip()):
                standalone_blocks.append((p, match.group(1)))
            else:
                inline_paragraphs.append(p)

        for p, name in standalone_blocks:
            markdown_text = block_content.get(name, "")
            if markdown_text.strip():
                _insert_block_content_after(p, markdown_text, output_dir)
            # Remove the marker paragraph; guard against detached elements
            parent = p._p.getparent()
            if parent is not None:
                parent.remove(p._p)
            else:
                # Fallback: clear the paragraph text so the marker disappears
                for run in p.runs:
                    run.text = ""

        for p in inline_paragraphs:
            _replace_inline_blocks_in_paragraph(p, block_content)

        doc.save(output_path)
        print(f"[DOC_GEN] render_placeholder_docx: saved filled DOCX to {output_path} "
              f"({len(values)} placeholder(s), {len(standalone_blocks)} standalone block(s), "
              f"{len(inline_paragraphs)} inline block paragraph(s))")
    except Exception as e:
        print(f"[DOC_GEN] render_placeholder_docx: FAILED — {e}")

    return {"final_output": f"[DOC_GEN] saved DOCX to {output_path}"}
